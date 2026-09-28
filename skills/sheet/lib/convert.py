"""文件格式转换：统一用 xlsx 作为工作副本。

支持 .xls / .xlsx / .csv / .ods。转换走 LibreOffice（已预装）。
"""
import csv
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from . import UserError


SUPPORTED_EXTS = {'.xls', '.xlsx', '.csv', '.ods'}
_CSV_ENCODINGS = ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'latin-1')


def check_supported(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise UserError(
            f"不支持的格式: {ext}（支持: {', '.join(sorted(SUPPORTED_EXTS))}）"
        )
    return ext


def to_workbook(path: Path, work_dir: Path) -> Path:
    """把输入文件转成 xlsx 工作副本，返回 xlsx 路径。"""
    ext = check_supported(path)
    work_dir.mkdir(parents=True, exist_ok=True)

    if ext == '.xlsx':
        dst = work_dir / path.name
        shutil.copy2(path, dst)
        return dst

    if ext == '.csv':
        rows = _read_csv(path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        # CSV 是文本格式：原样保留字段字符串（不做数字 coerce），回写才算真「最小改动」
        for row in rows:
            ws.append(list(row))
        dst = work_dir / (path.stem + '.xlsx')
        wb.save(dst)
        return dst

    return _libreoffice(path, work_dir, '.xlsx')


def from_workbook(xlsx: Path, dst: Path, original_ext: str):
    """把处理后的 xlsx 转回目标格式，写到 dst。"""
    ext = original_ext.lower()
    dst.parent.mkdir(parents=True, exist_ok=True)

    if ext == '.xlsx':
        shutil.copy2(xlsx, dst)
        return

    if ext == '.csv':
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        ws = wb.active
        with open(dst, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(['' if v is None else v for v in row])
        return

    with tempfile.TemporaryDirectory() as td:
        produced = _libreoffice(xlsx, Path(td), ext)
        shutil.copy2(produced, dst)


# ---------- 内部 ----------

def _lo_profile(out_dir: Path) -> Path:
    """LO 隔离 profile：执行容器根文件系统只读，默认 ~/.config 写入会直接 rc=77 崩（实测）；
    每次调用独立 profile 同时解决并发转换的 LO 单实例锁冲突。"""
    prof = out_dir / 'loprofile'
    (prof / 'user').mkdir(parents=True, exist_ok=True)
    return prof


def _libreoffice(src: Path, out_dir: Path, target_ext: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    target = target_ext.lstrip('.')
    r = subprocess.run(
        ['libreoffice', '--headless', '--norestore',
         '-env:UserInstallation=' + _lo_profile(out_dir).as_uri(),
         '--convert-to', target,
         '--outdir', str(out_dir), str(src)],
        capture_output=True, text=True, timeout=180,
    )
    if r.returncode != 0:
        msg = (r.stderr or r.stdout).strip() or '无输出'
        raise UserError(f"LibreOffice 转换失败: {msg}")
    expected = out_dir / (src.stem + target_ext)
    if not expected.exists():
        raise UserError(f"LibreOffice 未生成输出: {expected}")
    return expected


# 重算适用的二进制/文档格式（csv 无缓存概念，跳过）
RECALC_EXTS = {'.xlsx', '.xls', '.ods'}


def recalc(path: Path, work_dir: Path):
    """LibreOffice 强制重算公式缓存值：openpyxl 产物无缓存，LO 加载时必须求值才能保存，
    实测（LO 7.4）重算结果会写入缓存元素且公式文本保留。

    返回 (重算产物路径, None) 或 (None, 失败原因)；不抛异常，是否回退由调用方决策。
    注意返回的文件在 work_dir 内，调用方需在同一生命周期内用完。"""
    ext = path.suffix.lower()
    if ext not in RECALC_EXTS:
        return None, f"{ext} 格式无需重算"
    out_dir = work_dir / 'recalc'
    out_dir.mkdir(parents=True, exist_ok=True)
    expected = out_dir / path.name
    try:
        r = subprocess.run(
            ['libreoffice', '--headless', '--norestore',
             '-env:UserInstallation=' + _lo_profile(out_dir).as_uri(),
             '--convert-to', ext[1:],
             '--outdir', str(out_dir), str(path)],
            capture_output=True, text=True, timeout=180,
        )
    except FileNotFoundError:
        return None, "LibreOffice 不可用"
    except subprocess.TimeoutExpired:
        return None, "LibreOffice 重算超时(180s)"
    if r.returncode != 0 or not expected.exists():
        msg = (r.stderr or r.stdout).strip()[:150] or '无输出'
        return None, f"LibreOffice 重算失败: {msg}"
    return expected, None


def _read_csv(path: Path):
    last_err = None
    for enc in _CSV_ENCODINGS:
        try:
            with open(path, encoding=enc, newline='') as f:
                return list(csv.reader(f))
        except UnicodeDecodeError as e:
            last_err = e
            continue
    raise UserError(f"无法识别 CSV 编码: {path}（最后错误: {last_err}）")
