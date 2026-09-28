"""sheet-edit 的各个 op 实现。"""
import datetime
import shutil
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.formula import ArrayFormula

from . import UserError, convert, formulas, locate

# inspect 返回的公式清单条数上限（大表防爆 token）
MAX_FORMULAS = 500


def _json_safe(v):
    """单元格值 → JSON 可序列化形态。"""
    if isinstance(v, ArrayFormula):
        return {'formula': v.text}
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S') if (v.hour or v.minute or v.second) else v.strftime('%Y-%m-%d')
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return str(v)
    return v


def _cell_text(v):
    """公式统一成字符串形态（ArrayFormula → 其文本），供内部比较用。"""
    return v.text if isinstance(v, ArrayFormula) else v


def _is_formula_val(v):
    return locate.is_formula(v) or isinstance(v, ArrayFormula)


# ============ inspect ============

def inspect(args):
    file = _req_file(args)
    header_row = _int(args, 'header_row', 1)
    limit = _int(args, 'limit', 50)
    sheet_filter = args.get('sheet')

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)

        sheets = []
        found = False
        for name in wb.sheetnames:
            if sheet_filter and name != sheet_filter:
                continue
            found = True
            ws = wb[name]
            sheets.append(_sheet_info(ws, name, header_row, limit))

        if sheet_filter and not found:
            raise UserError(
                f"sheet 不存在: {sheet_filter}（可用: {', '.join(wb.sheetnames)}）"
            )

    return {'ok': True, 'sheets': sheets}


def _sheet_info(ws, name, header_row, limit):
    info = {
        'name': name,
        'rows': ws.max_row,
        'cols': ws.max_column,
        'header_row': header_row,
        'suggested_header_row': locate.detect_header_row(ws),
        'headers': [
            _json_safe(ws.cell(header_row, c).value)
            for c in range(1, ws.max_column + 1)
        ],
        'preview': [],
        'formulas': [],
        'formulas_truncated': False,
        'merged': [str(r) for r in ws.merged_cells.ranges],
        'truncated': False,
    }
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if _is_formula_val(v):
                if len(info['formulas']) < MAX_FORMULAS:
                    info['formulas'].append({'cell': cell.coordinate, 'formula': _cell_text(v)})
                else:
                    info['formulas_truncated'] = True
    end = min(header_row + limit, ws.max_row)
    for r in range(header_row + 1, end + 1):
        row_data = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            row_data.append({'formula': _cell_text(v)} if _is_formula_val(v) else _json_safe(v))
        info['preview'].append(row_data)
    info['truncated'] = ws.max_row > end
    return info


# ============ apply ============

def apply(args):
    file = _req_file(args)
    header_row = _int(args, 'header_row', 1)
    changes = args.get('changes')
    if not isinstance(changes, list) or not changes:
        raise UserError("changes 必须是非空数组")

    original_ext = convert.check_supported(file)
    out_path = _out_path(file, args.get('output'))

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)

        plan = []
        for ch in changes:
            plan.extend(_resolve_change(wb, ch, header_row))

        for item in plan:
            if not locate.value_eq(item['expected_old'], item['actual_old']):
                raise UserError(
                    f"旧值不匹配: {item['sheet']} 第 {item['row']} 行 "
                    f"'{item['col_name']}': 期望 {item['expected_old']!r}, "
                    f"实际 {_json_safe(item['actual_old'])!r}"
                )

        before = _snapshot(wb)

        for item in plan:
            ws = wb[item['sheet']]
            _assert_writable(ws, item['row'], item['col'])
            ws.cell(item['row'], item['col']).value = item['new']

        _save(wb, xlsx)
        convert.from_workbook(xlsx, out_path, original_ext)
        try:
            _verify_apply(out_path, before, plan)
        except Exception:
            _drop(out_path)
            raise
        recalc_note = _try_recalc(out_path, Path(td), lambda p: _verify_apply(p, before, plan))

    return {
        'ok': True,
        'output': out_path.name,
        'recalc': recalc_note,
        'changes': [
            {
                'sheet': it['sheet'],
                'row': it['row'],
                'col_name': it['col_name'],
                'old': _json_safe(it['actual_old']),
                'new': _json_safe(it['new']),
                'match': it['match'],
            }
            for it in plan
        ],
    }


def _resolve_change(wb, ch, header_row):
    sheet = ch.get('sheet')
    if not sheet:
        raise UserError("change 缺少 sheet")
    if sheet not in wb.sheetnames:
        raise UserError(
            f"sheet 不存在: {sheet}（可用: {', '.join(wb.sheetnames)}）"
        )
    match = ch.get('match') or {}
    set_spec = ch.get('set') or {}
    if not isinstance(set_spec, dict) or not set_spec:
        raise UserError("change.set 必须是非空对象")

    ws = wb[sheet]
    col_map = locate.build_col_map(ws, header_row)
    row = locate.find_row(ws, col_map, match, header_row)

    items = []
    for col_name, spec in set_spec.items():
        if not isinstance(spec, dict) or 'old' not in spec or 'new' not in spec:
            raise UserError(
                f"set['{col_name}'] 必须是 {{\"old\": ..., \"new\": ...}} 格式"
            )
        col = locate.find_col(col_map, col_name)
        actual = ws.cell(row, col).value
        if isinstance(actual, ArrayFormula):
            raise UserError(
                f"{sheet}!{get_column_letter(col)}{row} 是数组公式，"
                "本工具不支持修改（请在 Excel/WPS 手工处理）"
            )
        new = spec['new']
        if isinstance(new, (dict, list)):
            raise UserError(f"set['{col_name}'].new 只能是数字/字符串/布尔/null，收到 {new!r}")
        items.append({
            'sheet': sheet,
            'row': row,
            'col': col,
            'col_name': col_name,
            'match': dict(match),
            'expected_old': spec['old'],
            'actual_old': actual,
            'new': new,
        })
    return items


def _verify_apply(out_path, before, plan):
    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(out_path, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)

        for item in plan:
            v = wb[item['sheet']].cell(item['row'], item['col']).value
            if not locate.value_eq(v, item['new']):
                raise UserError(
                    f"校验失败: {item['sheet']} 第 {item['row']} 行 "
                    f"'{item['col_name']}' 应为 {item['new']!r}, 实际 {v!r}"
                )

        targets = {(it['sheet'], it['row'], it['col']) for it in plan}
        diffs = []
        for name in wb.sheetnames:
            ws = wb[name]
            before_sheet = before.get(name, {})
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    if (name, r, c) in targets:
                        continue
                    v = _cell_text(ws.cell(r, c).value)
                    if locate.is_formula(v):
                        continue  # apply 不动公式；公式若有变化在下面单独盯
                    if (r, c) not in before_sheet:
                        if _notempty(v):
                            diffs.append(f"{name} ({r},{c}): 新增 {v!r}")
                        continue
                    if not locate.value_eq(before_sheet[(r, c)], v):
                        diffs.append(
                            f"{name} ({r,c}): {before_sheet[(r,c)]!r} → {v!r}"
                        )
            # 公式原文必须一字不变（apply 语义：只改值，不动公式）
            for (r, c), v0 in before_sheet.items():
                if not locate.is_formula(v0):
                    continue
                if (name, r, c) in targets:
                    continue
                v1 = _cell_text(ws.cell(r, c).value)
                if v1 != v0:
                    diffs.append(f"{name} 公式被改动 ({r},{c}): {v0!r} → {v1!r}")
        if diffs:
            raise UserError(
                "校验失败: 非目标单元格发生变化:\n  " + "\n  ".join(diffs[:20])
            )


# ============ insert_rows / delete_rows ============

def insert_rows(args):
    file = _req_file(args)
    sheet = _req_sheet(args)
    count = _pos_int(args, 'count', 1)
    header_row = _int(args, 'header_row', 1)
    fill = args.get('fill') or {}
    after_match = args.get('after_match') or {}
    if fill and not isinstance(fill, dict):
        raise UserError("fill 必须是对象")

    original_ext = convert.check_supported(file)
    out_path = _out_path(file, args.get('output'))

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        _check_sheet(wb, sheet)

        ws = wb[sheet]
        col_map = locate.build_col_map(ws, header_row)
        anchor = locate.find_row(ws, col_map, after_match, header_row)
        insert_at = anchor + 1

        fill_cells = {}
        for col_name, value in fill.items():
            fill_cells[col_name] = (locate.find_col(col_map, col_name), value)

        before = _snapshot(wb)
        _adjust_formulas(wb, sheet, 'row', insert_at, count, delete=False)
        ws.insert_rows(insert_at, count)
        _shift_merged(ws, 'row', insert_at, count, delete=False)
        _materialize(ws, insert_at, insert_at + count - 1, 'row')
        expected_new = {}
        for col_name, (c, value) in fill_cells.items():
            for r in range(insert_at, insert_at + count):
                _assert_writable(ws, r, c, where=f"fill 列 '{col_name}' 新行")
                ws.cell(r, c).value = value
                expected_new[(r, c)] = value

        _save(wb, xlsx)
        convert.from_workbook(xlsx, out_path, original_ext)
        try:
            _verify_struct(out_path, before, sheet, 'row', insert_at, count, False, expected_new)
        except Exception:
            _drop(out_path)
            raise
        recalc_note = _try_recalc(out_path, Path(td),
                                  lambda p: _verify_struct(p, before, sheet, 'row', insert_at, count, False, expected_new))

    return {
        'ok': True,
        'output': out_path.name,
        'recalc': recalc_note,
        'sheet': sheet,
        'inserted_at': insert_at,
        'count': count,
    }


def delete_rows(args):
    file = _req_file(args)
    sheet = _req_sheet(args)
    count = _pos_int(args, 'count', 1)
    header_row = _int(args, 'header_row', 1)
    match = args.get('match') or {}

    original_ext = convert.check_supported(file)
    out_path = _out_path(file, args.get('output'))

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        _check_sheet(wb, sheet)

        ws = wb[sheet]
        col_map = locate.build_col_map(ws, header_row)
        start = locate.find_row(ws, col_map, match, header_row)

        before = _snapshot(wb)
        _adjust_formulas(wb, sheet, 'row', start, count, delete=True)
        ws.delete_rows(start, count)
        _shift_merged(ws, 'row', start, count, delete=True)

        _save(wb, xlsx)
        convert.from_workbook(xlsx, out_path, original_ext)
        try:
            _verify_struct(out_path, before, sheet, 'row', start, count, True, {})
        except Exception:
            _drop(out_path)
            raise
        recalc_note = _try_recalc(out_path, Path(td),
                                  lambda p: _verify_struct(p, before, sheet, 'row', start, count, True, {}))

    return {
        'ok': True,
        'output': out_path.name,
        'recalc': recalc_note,
        'sheet': sheet,
        'deleted_at': start,
        'count': count,
    }


# ============ insert_cols / delete_cols ============

def insert_cols(args):
    file = _req_file(args)
    sheet = _req_sheet(args)
    after = args.get('after')
    if not after:
        raise UserError("缺少 after（在此列之后插入）")
    count = _pos_int(args, 'count', 1)
    header_row = _int(args, 'header_row', 1)
    fill_header = args.get('fill_header')
    if fill_header and count > 1:
        raise UserError("fill_header 目前只支持 count=1")

    original_ext = convert.check_supported(file)
    out_path = _out_path(file, args.get('output'))

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        _check_sheet(wb, sheet)

        ws = wb[sheet]
        col_map = locate.build_col_map(ws, header_row)
        anchor = locate.find_col(col_map, after)
        insert_at = anchor + 1

        before = _snapshot(wb)
        _adjust_formulas(wb, sheet, 'col', insert_at, count, delete=False)
        ws.insert_cols(insert_at, count)
        _shift_merged(ws, 'col', insert_at, count, delete=False)
        _materialize(ws, insert_at, insert_at + count - 1, 'col')
        expected_new = {}
        if fill_header:
            _assert_writable(ws, header_row, insert_at, where="fill_header")
            ws.cell(header_row, insert_at).value = fill_header
            expected_new[(header_row, insert_at)] = fill_header

        _save(wb, xlsx)
        convert.from_workbook(xlsx, out_path, original_ext)
        try:
            _verify_struct(out_path, before, sheet, 'col', insert_at, count, False, expected_new)
        except Exception:
            _drop(out_path)
            raise
        recalc_note = _try_recalc(out_path, Path(td),
                                  lambda p: _verify_struct(p, before, sheet, 'col', insert_at, count, False, expected_new))

    return {
        'ok': True,
        'output': out_path.name,
        'recalc': recalc_note,
        'sheet': sheet,
        'inserted_at': insert_at,
        'count': count,
    }


def delete_cols(args):
    file = _req_file(args)
    sheet = _req_sheet(args)
    col_name = args.get('col')
    if not col_name:
        raise UserError("缺少 col（要删除的列名）")
    count = _pos_int(args, 'count', 1)
    header_row = _int(args, 'header_row', 1)

    original_ext = convert.check_supported(file)
    out_path = _out_path(file, args.get('output'))

    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(file, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        _check_sheet(wb, sheet)

        ws = wb[sheet]
        col_map = locate.build_col_map(ws, header_row)
        start = locate.find_col(col_map, col_name)

        before = _snapshot(wb)
        _adjust_formulas(wb, sheet, 'col', start, count, delete=True)
        ws.delete_cols(start, count)
        _shift_merged(ws, 'col', start, count, delete=True)

        _save(wb, xlsx)
        convert.from_workbook(xlsx, out_path, original_ext)
        try:
            _verify_struct(out_path, before, sheet, 'col', start, count, True, {})
        except Exception:
            _drop(out_path)
            raise
        recalc_note = _try_recalc(out_path, Path(td),
                                  lambda p: _verify_struct(p, before, sheet, 'col', start, count, True, {}))

    return {
        'ok': True,
        'output': out_path.name,
        'recalc': recalc_note,
        'sheet': sheet,
        'deleted_at': start,
        'count': count,
    }


# ============ 内部工具 ============

def _req_file(args) -> Path:
    f = args.get('file')
    if not f:
        raise UserError("缺少 file 参数")
    p = Path(f)
    if not p.exists():
        raise UserError(f"文件不存在: {f}")
    return p


def _req_sheet(args) -> str:
    s = args.get('sheet')
    if not s:
        raise UserError("缺少 sheet 参数")
    return s


def _check_sheet(wb, sheet):
    if sheet not in wb.sheetnames:
        raise UserError(
            f"sheet 不存在: {sheet}（可用: {', '.join(wb.sheetnames)}）"
        )


def _int(args, key, default):
    v = args.get(key)
    if v is None:
        return default
    try:
        return int(v)
    except (TypeError, ValueError):
        raise UserError(f"{key} 必须是整数: {v!r}")


def _pos_int(args, key, default):
    v = _int(args, key, default)
    if v < 1:
        raise UserError(f"{key} 必须 >= 1: {v}")
    return v


def _out_path(file: Path, output):
    if output:
        name = str(output)
        if name in ('.', '..'):
            raise UserError(f"output 无效: {name!r}")
        p = Path(name)
        # 绝对路径原样用：skills.js 会把 output 这类路径参数预处理成会话目录下的绝对路径，
        # 拒含分隔符反而把平台惯例变成必死参数；防越权由执行容器自身的会话隔离担责。
        return p if p.is_absolute() else file.parent / p
    return file.parent / f"{file.stem}（已改）{file.suffix}"


def _assert_writable(ws, r, c, where=None):
    """目标格落在合并区域内且不是锚点 → 拒写（openpyxl 写 MergedCell 会裸 AttributeError）。"""
    for cr in ws.merged_cells.ranges:
        if cr.min_row <= r <= cr.max_row and cr.min_col <= c <= cr.max_col:
            if (r, c) != (cr.min_row, cr.min_col):
                loc = f"{ws.title}!{get_column_letter(c)}{r}"
                raise UserError(
                    f"{where or loc} 落在合并区域 {cr.coord} 内且不是锚点格，"
                    "请改写到该区域左上角锚点"
                )
            return


def _save(wb, xlsx):
    # openpyxl 保存会丢公式缓存值；置 fullCalcOnLoad 让 Excel/WPS/LibreOffice 打开时全量重算
    wb.calculation.fullCalcOnLoad = True
    wb.save(xlsx)


def _drop(path: Path):
    """交付校验失败时移除坏产物，兑现「不交付」。"""
    try:
        path.unlink()
    except OSError:
        pass


def _try_recalc(out_path, work_dir, verify):
    """正确性门禁通过后尝试 LO 重算；重算产物同样过一遍 verify 才能覆盖交付，
    否则保持未重算版（仍带 fullCalcOnLoad 标记，Excel/WPS 打开会自动算）。
    重算是可用性增强不是正确性条件，故失败只降级不报错。返回状态供结果 JSON。"""
    cand, err = convert.recalc(out_path, work_dir)
    if cand is None:
        return 'skipped' if '无需重算' in (err or '') else f'failed: {err}'
    try:
        verify(cand)
    except Exception as e:
        return f'failed: 重算产物未过自检，交付未重算版（{str(e)[:120]}）'
    shutil.copy2(cand, out_path)
    return 'done'


def _adjust_formulas(wb, target_sheet, axis, at, count, delete):
    """全工作簿调整公式：含「其他表引用目标表」的跨表公式；目标表引用他表的不动（同 Excel）。

    数组公式/外部簿引用/结构化引用 → UserError，整体放弃（调用方删产物）。
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, ArrayFormula):
                    raise UserError(
                        f"不支持数组公式（{ws.title}!{cell.coordinate}），"
                        "请先在 Excel/WPS 中将其转为普通公式后重试"
                    )
                if locate.is_formula(v):
                    cell.value = formulas.shift_formula(
                        v, ws.title, target_sheet, axis, at, count, delete)


def _shift_merged(ws, axis, at, count, delete):
    """合并区域随插删移动（openpyxl 原生不管 merged_cells）。

    插入：边界 >= at 后移 → 跨界插入时区域自动扩大（同 Excel）。
    删除：区域整体落在删除带内 → 收缩到带前；完全被删 → 消失。
           部分重叠（一头在带内、另一头在外）Excel 语义含糊 → 直接拒绝。
    """
    def fully_in(v):
        return at <= v < at + count

    def mv(v):
        if not delete:
            return v + count if v >= at else v
        if v < at:
            return v
        if v < at + count:
            return at - 1
        return v - count

    moved = []
    for cr in ws.merged_cells.ranges:
        r1, r2, c1, c2 = cr.min_row, cr.max_row, cr.min_col, cr.max_col
        if delete:
            if axis == 'row' and fully_in(r1) != fully_in(r2):
                raise UserError(f"不支持删除跨越合并区域 {cr.coord} 边界的行（Excel 行为无法推断）")
            if axis == 'col' and fully_in(c1) != fully_in(c2):
                raise UserError(f"不支持删除跨越合并区域 {cr.coord} 边界的列（Excel 行为无法推断）")
        nr1, nr2, nc1, nc2 = (mv(r1), mv(r2), c1, c2) if axis == 'row' else (r1, r2, mv(c1), mv(c2))
        if nr1 == nr2 and nc1 == nc2:
            continue  # 区域整体消失
        moved.append(str(CellRange(min_col=nc1, min_row=nr1, max_col=nc2, max_row=nr2)))
    ws.merged_cells = MultiCellRange(ranges=set(moved))


def _materialize(ws, a, b, axis):
    """把新插的空行/空列注册进单元格表（否则 openpyxl 不扩大 dimension，csv 回写会丢行）。"""
    if axis == 'row':
        for r in range(a, b + 1):
            ws.cell(r, 1)
    else:
        for c in range(a, b + 1):
            ws.cell(1, c)


def _snapshot(wb):
    """记录所有单元格（含公式原文），供插删后的结构化自检。"""
    snap = {}
    for ws in wb.worksheets:
        d = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    d[(cell.row, cell.column)] = _cell_text(cell.value)
        snap[ws.title] = d
    return snap


def _map_axis(v, at, count, delete):
    """旧坐标 → 新坐标；None 表示落在被删除带内。"""
    if not delete:
        return v + count if v >= at else v
    if v < at:
        return v
    if v < at + count:
        return None
    return v - count


def _notempty(v):
    return v is not None and v != ''


def _cmp_norm(v):
    """比较归一：csv 往返的空串视同空。"""
    return None if v == '' else v


def _verify_struct(out_path, before, sheet, axis, at, count, delete, expected_new):
    """插删自检：目标表尺寸正确；旧格按坐标映射逐一存在（公式等于预期调整结果）；
    非操作表尺寸不变、内容仅允许「引用目标表的公式」发生调整；新带区域除 fill 外为空。"""
    with tempfile.TemporaryDirectory() as td:
        xlsx = convert.to_workbook(out_path, Path(td))
        wb = openpyxl.load_workbook(xlsx, data_only=False)

        diffs = []
        for ws in wb.worksheets:
            is_target = ws.title == sheet
            bf = before.get(ws.title, {})
            rows_b = max((r for r, _ in bf), default=0)
            cols_b = max((c for _, c in bf), default=0)

            if is_target:
                exp_rows = _map_bounds(rows_b, axis, 'row', at, count, delete)
                exp_cols = _map_bounds(cols_b, axis, 'col', at, count, delete)
                if ws.max_row != exp_rows:
                    diffs.append(f"{ws.title}: 行数应为 {exp_rows} 实际 {ws.max_row}")
                if ws.max_column != exp_cols:
                    diffs.append(f"{ws.title}: 列数应为 {exp_cols} 实际 {ws.max_column}")
            else:
                if ws.max_row != rows_b or ws.max_column != cols_b:
                    diffs.append(
                        f"{ws.title}: 非操作表尺寸被波及 {rows_b}x{cols_b} → {ws.max_row}x{ws.max_column}")

            for (r, c), v in bf.items():
                nr, nc = r, c
                if is_target:
                    nr = _map_axis(r, at, count, delete) if axis == 'row' else r
                    nc = _map_axis(c, at, count, delete) if axis == 'col' else c
                    if nr is None or nc is None:
                        continue
                actual = _cmp_norm(_cell_text(ws.cell(nr, nc).value))
                expected = v
                if locate.is_formula(v):
                    expected = formulas.shift_formula(v, ws.title, sheet, axis, at, count, delete)
                if not locate.value_eq(_cmp_norm(expected), actual):
                    diffs.append(f"{ws.title} ({r},{c})→({nr},{nc}): 期望 {expected!r} 实际 {actual!r}")

        if not delete:
            ws = wb[sheet]
            band = range(at, at + count)
            rows = band if axis == 'row' else range(1, ws.max_row + 1)
            cols = range(1, ws.max_column + 1) if axis == 'row' else band
            for r in rows:
                for c in cols:
                    v = _cmp_norm(_cell_text(ws.cell(r, c).value))
                    if (r, c) in expected_new:
                        if not locate.value_eq(_cmp_norm(expected_new[(r, c)]), v):
                            diffs.append(f"{sheet} fill({r},{c}) 应为 {expected_new[(r, c)]!r} 实际 {v!r}")
                    elif _notempty(v):
                        diffs.append(f"{sheet} 新区域({r},{c}) 应为空 实际 {v!r}")

        if diffs:
            raise UserError("校验失败: 插删后结构与预期不符:\n  " + "\n  ".join(diffs[:20]))


def _map_bounds(n, _axis, kind, at, count, delete):
    """快照尺寸 → 预期新尺寸（按轴）。"""
    if kind != _axis:
        return n
    if not delete:
        return n + count
    if n < at:
        return n
    return n - min(count, n - at + 1)
